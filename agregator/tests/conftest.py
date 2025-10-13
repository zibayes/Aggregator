import pytest
from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.contrib.auth import get_user_model
import tempfile
from django.utils import timezone
from datetime import timedelta
from django_celery_results.models import TaskResult
import jwt
import os
from unittest.mock import MagicMock

User = get_user_model()


@pytest.fixture
def client():
    from django.test import Client
    return Client()


@pytest.fixture
def api_client():
    from rest_framework.test import APIClient
    return APIClient()


@pytest.fixture
def test_user(db):
    """Создает основного тестового пользователя"""
    return User.objects.create_user(
        username='testuser',
        password='testpass123',
        email='test@example.com',
        first_name='Test',
        last_name='User'
    )


@pytest.fixture
def test_user_2(db):
    """Создает второго тестового пользователя"""
    return User.objects.create_user(
        username='otheruser',
        password='otherpass123',
        email='other@example.com'
    )


@pytest.fixture
def admin_user(db):
    return User.objects.create_superuser(
        username='admin',
        password='adminpass123',
        email='admin@example.com'
    )


@pytest.fixture
def avatar_file():
    """Создает тестовый файл аватара"""
    return SimpleUploadedFile(
        "avatar.jpg",
        b"file_content",
        content_type="image/jpeg"
    )


@pytest.fixture
def test_image():
    return SimpleUploadedFile(
        "test_avatar.jpg",
        b"file_content",
        content_type="image/jpeg"
    )


@pytest.fixture
def test_act(db, test_user):
    from agregator.models import Act
    return Act.objects.create(
        user=test_user,
        year='2023',
        name_number='Test Act',
        type='ГИКЭ',
        place='Test Location'
    )


@pytest.fixture
def test_scientific_report(db, test_user):
    from agregator.models import ScientificReport
    return ScientificReport.objects.create(
        user=test_user,
        name='Test Scientific Report',
        organization='Test Org',
        author='Test Author'
    )


@pytest.fixture
def test_tech_report(db, test_user):
    from agregator.models import TechReport
    return TechReport.objects.create(
        user=test_user,
        name='Test Tech Report',
        organization='Test Org',
        author='Test Author'
    )


@pytest.fixture
def test_open_list(db, test_user):
    from agregator.models import OpenLists
    return OpenLists.objects.create(
        user=test_user,
        number='TEST-001',
        holder='Test Holder',
        object='Test Object'
    )


@pytest.fixture
def test_archaeological_heritage_site(db):
    from agregator.models import ArchaeologicalHeritageSite
    return ArchaeologicalHeritageSite.objects.create(
        doc_name='Test OAN',
        district='Test District',
        register_num='TEST-OAN-001'
    )


@pytest.fixture
def test_identified_heritage_site(db):
    from agregator.models import IdentifiedArchaeologicalHeritageSite
    return IdentifiedArchaeologicalHeritageSite.objects.create(
        name='Test VOAN',
        address='Test Address',
        obj_info='Test Info'
    )


@pytest.fixture
def test_account_card(db, test_user, test_identified_heritage_site):
    from agregator.models import ObjectAccountCard
    card = ObjectAccountCard.objects.create(
        user=test_user,
        name='Test Account Card',
        creation_time='Test Period',
        address='Test Address'
    )
    # Связываем с heritage site
    test_identified_heritage_site.account_card = card
    test_identified_heritage_site.save()
    return card


@pytest.fixture
def test_commercial_offer(test_user):
    from agregator.models import CommercialOffers
    return CommercialOffers.objects.create(
        user=test_user,
        origin_filename='Test Commercial Offer'
    )


@pytest.fixture
def test_geo_object(test_user):
    from agregator.models import GeoObject
    return GeoObject.objects.create(
        user=test_user,
        origin_filename='Test Geo Object'
    )


@pytest.fixture
def test_tasks(db, test_user):
    """Создает тестовые задачи для пользователя"""
    from agregator.models import UserTasks

    tasks = []
    for i, task_type in enumerate(
            ['act', 'scientific_report', 'tech_report', 'open_list', 'account_card', 'commercial_offer', 'geo_object']):
        task = UserTasks.objects.create(
            user=test_user,
            task_id=f'test_task_{i}',
            files_type=task_type,
            upload_source={'source': 'Пользовательский файл'}
        )
        tasks.append(task)

    # Создаем TaskResult для этих задач
    for task in tasks:
        TaskResult.objects.create(
            task_id=task.task_id,
            status='SUCCESS',
            result='{"message": "success"}'
        )

    return tasks


@pytest.fixture
def admin_tasks(db, admin_user):
    """Создает тестовые задачи для админа (external source)"""
    from agregator.models import UserTasks

    tasks = []
    for i, task_type in enumerate(['act', 'scientific_report', 'tech_report', 'open_list']):
        task = UserTasks.objects.create(
            user=admin_user,
            task_id=f'admin_task_{i}',
            files_type=task_type,
            upload_source={'source': 'external'}
        )
        tasks.append(task)

    for task in tasks:
        TaskResult.objects.create(
            task_id=task.task_id,
            status='SUCCESS',
            result='{"message": "success"}'
        )

    return tasks


@pytest.fixture
def wopi_test_file():
    """Создает тестовый файл для WOPI"""
    with tempfile.NamedTemporaryFile(delete=False, suffix='.txt') as f:
        f.write(b'Test file content for WOPI')
        return f.name


@pytest.fixture
def wopi_token(test_user, wopi_test_file):
    """Генерирует валидный WOPI токен для конкретного тестового файла"""
    from agregator.wopi.views import generate_wopi_token
    file_id = os.path.basename(wopi_test_file)
    return generate_wopi_token(
        user_id=test_user.id,
        username=test_user.username,
        file_path=file_id,
        can_write=True
    )


@pytest.fixture
def wopi_token_readonly(test_user):
    """Генерирует WOPI токен только для чтения"""
    from agregator.wopi.views import generate_wopi_token
    return generate_wopi_token(
        user_id=test_user.id,
        username=test_user.username,
        file_path='test_file.txt',
        can_write=False
    )


@pytest.fixture
def wopi_token_expired(test_user):
    """Генерирует просроченный WOPI токен"""
    from agregator.wopi.views import WOPI_ACCESS_SECRET
    payload = {
        'user_id': test_user.id,
        'username': test_user.username,
        'file_path': 'test_file.txt',
        'can_write': True,
        'exp': timezone.now() - timedelta(hours=1),
        'ttl': 3600
    }
    return jwt.encode(payload, WOPI_ACCESS_SECRET, algorithm='HS256')


@pytest.fixture
def wopi_invalid_token():
    """Генерирует невалидный токен"""
    return 'invalid.token.here'


@pytest.fixture
def valid_pdf_file():
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    import io

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    c.drawString(100, 750, "Test PDF with coordinates")
    c.save()
    buffer.seek(0)
    return SimpleUploadedFile("test.pdf", buffer.read(), content_type='application/pdf')


@pytest.fixture
def valid_docx_file():
    from docx import Document
    import io

    doc = Document()
    doc.add_paragraph("Test DOCX with coordinates")
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile("test.docx", buffer.read(),
                              content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


@pytest.fixture
def valid_xlsx_file():
    from openpyxl import Workbook
    import io

    wb = Workbook()
    ws = wb.active
    ws['A1'] = 'Test XLSX with coordinates'
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile("test.xlsx", buffer.read(),
                              content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
